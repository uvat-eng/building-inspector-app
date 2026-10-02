"""Сборка акта осмотра в формате Word (.docx).

Раньше акт отдавался как HTML с расширением .doc — телефон считает такой
файл битым и открывать отказывается. Здесь собирается настоящий Word.
"""

import io

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import tenant

ACCENT = RGBColor(0x1F, 0x3A, 0x6E)
GREY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xB3, 0x26, 0x1E)

MONTHS = [
    'января', 'февраля', 'марта', 'апреля', 'мая', 'июня',
    'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря',
]


def ru_date(value):
    """Дата словами: 30 сентября 2026 г."""
    if not value:
        return '—'
    text = str(value)[:10]
    try:
        if '.' in text:
            day, month, year = text.split('.')
        else:
            year, month, day = text.split('-')
        return f'{int(day)} {MONTHS[int(month) - 1]} {year} г.'
    except Exception:
        return text


def shade(cell, color):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(el)


def set_font(doc):
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    fmt = style.paragraph_format
    fmt.space_after = Pt(4)
    fmt.line_spacing = 1.15


def para(doc, text='', size=12, bold=False, italic=False, color=None,
         align=None, space_before=0, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return p


def kv_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for name, value in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(5.5)
        cells[1].width = Cm(11.5)
        shade(cells[0], 'F2F4F8')
        for cell, text, bold in ((cells[0], name, False), (cells[1], value or '—', True)):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(text))
            run.bold = bold
            run.font.size = Pt(11)
    return table


def defects_table(doc, defects):
    """Таблица выявленных замечаний."""
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    titles = ('№', 'Выявленное замечание', 'Норматив', 'Значимость', 'Устранить до')
    widths = (Cm(0.9), Cm(7.6), Cm(4.4), Cm(2.2), Cm(2.4))
    for cell, title, width in zip(table.rows[0].cells, titles, widths):
        cell.width = width
        shade(cell, 'E8ECF4')
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(9.5)

    for i, d in enumerate(defects, 1):
        cells = table.add_row().cells
        values = (
            str(i),
            str(d.get('title') or '—'),
            str(d.get('normRef') or 'норматив не указан'),
            str(d.get('severity') or '—'),
            str(d.get('deadline') or '—'),
        )
        for cell, text, width in zip(cells, values, widths):
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.size = Pt(9.5)
    return table


def sign_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for name, who in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(8.5)
        cells[1].width = Cm(8.5)
        for cell, text in ((cells[0], name), (cells[1], who)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(str(text))
            run.font.size = Pt(10)
    return table


def build_act_docx(insp, defects, object_title='', contractor=None):
    """Формирует .docx акта осмотра и возвращает его байты."""
    doc = Document()
    set_font(doc)

    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(1.4)

    defects = defects or []
    contractor = contractor or {}

    para(doc, tenant.org_name(), size=11, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, 'Служба строительного контроля', size=9.5, italic=True, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    para(doc, f"АКТ ОСМОТРА № {insp.get('number') or '—'}", size=15, bold=True,
         color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, 'о результатах строительного контроля', size=11, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    kv_table(doc, [
        ('Дата осмотра', ru_date(insp.get('createdAt'))),
        ('Объект', object_title),
        ('Вид работ', insp.get('workType')),
        ('Раздел проекта', insp.get('docRef')),
        ('Генподрядчик', insp.get('generalContractor') or contractor.get('name')),
        ('Субподрядчик', insp.get('subcontractor')),
        ('Представитель подрядчика', insp.get('contractorRep')),
        ('Осмотр провёл', insp.get('inspector')),
        ('Выявлено замечаний', str(len(defects))),
    ])

    para(doc, 'Выявленные замечания', size=12.5, bold=True, color=ACCENT,
         space_before=14, space_after=6)

    if defects:
        defects_table(doc, defects)
    else:
        para(doc, 'Замечаний не выявлено.', italic=True, color=GREY)

    note = str(insp.get('note') or '').strip()
    if note:
        para(doc, 'Примечание', size=12, bold=True, color=ACCENT,
             space_before=12, space_after=4)
        para(doc, note, size=11)

    if defects:
        para(doc,
             'Замечания подлежат устранению в указанные сроки. По результатам '
             'устранения проводится повторный осмотр.',
             size=10, italic=True, color=GREY, space_before=10, space_after=12)
    else:
        para(doc, 'Работы соответствуют проектной документации и требованиям '
                  'нормативных документов.',
             size=10, italic=True, color=GREY, space_before=10, space_after=12)

    sign_table(doc, [
        ('Осмотр провёл', insp.get('inspector') or ''),
        ('Представитель подрядчика', insp.get('contractorRep') or ''),
    ])

    para(doc, 'Подпись, расшифровка, дата', size=8.5, italic=True, color=GREY,
         space_before=2)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
