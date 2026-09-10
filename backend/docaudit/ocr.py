"""Распознавание сканов и фото документации через Cloud.ru."""

import base64
import io
import os
import re

import requests

CLOUD_URL = 'https://foundation-models.api.cloud.ru/v1/chat/completions'
OCR_MODEL = 'deepseek-ai/DeepSeek-OCR-2'
VL_MODEL = 'qwen/qwen3-vl-30b-a3b-instruct'

OCR_PROMPT = (
    'Извлеки ВЕСЬ текст с этого изображения документа построчно, сохраняя структуру. '
    'Таблицы передавай строками через |. Рукописные записи тоже распознай и помечай их '
    'как [от руки: текст]. Печати и штампы помечай как [печать: текст]. '
    'Подписи помечай как [подпись]. Не комментируй, выводи только распознанный текст.'
)


def cloud_key():
    return os.environ.get('CLOUDRU_API_KEY', '')


def ocr_image(raw: bytes, mime: str = 'image/jpeg', budget: int = 120) -> str:
    """Распознаёт одну страницу-картинку. Возвращает текст."""
    key = cloud_key()
    if not key:
        raise RuntimeError('Не задан ключ Cloud.ru — распознавание сканов недоступно')

    b64 = base64.b64encode(raw).decode()
    if not mime.startswith('image/'):
        mime = 'image/jpeg'

    last = ''
    for model in (VL_MODEL, OCR_MODEL):
        try:
            r = requests.post(
                CLOUD_URL,
                json={
                    'model': model,
                    'temperature': 0,
                    'max_tokens': 4000,
                    'messages': [
                        {
                            'role': 'user',
                            'content': [
                                {'type': 'text', 'text': OCR_PROMPT},
                                {
                                    'type': 'image_url',
                                    'image_url': {'url': f'data:{mime};base64,{b64}'},
                                },
                            ],
                        }
                    ],
                },
                headers={'Authorization': f'Bearer {key}'},
                timeout=budget,
            )
            r.raise_for_status()
            text = r.json()['choices'][0]['message']['content'] or ''
            if text.strip():
                return text.strip()
        except Exception as e:
            last = f'{type(e).__name__}: {str(e)[:200]}'
    if last:
        raise RuntimeError(f'Не удалось распознать страницу — {last}')
    return ''


def prep_image(raw: bytes) -> bytes:
    """Готовит фото к распознаванию: разворот по EXIF, выравнивание тона, сжатие."""
    try:
        from PIL import Image, ImageOps
    except Exception:
        return raw
    try:
        im = Image.open(io.BytesIO(raw))
        im = ImageOps.exif_transpose(im)
        if im.mode != 'RGB':
            im = im.convert('RGB')
        side = max(im.size)
        if side > 2200:
            k = 2200 / side
            im = im.resize((int(im.width * k), int(im.height * k)), Image.LANCZOS)
        im = ImageOps.autocontrast(im, cutoff=1)
        out = io.BytesIO()
        im.save(out, 'JPEG', quality=88)
        return out.getvalue()
    except Exception:
        return raw


def pdf_page_count(raw: bytes) -> int:
    try:
        import pypdf

        return len(pypdf.PdfReader(io.BytesIO(raw)).pages)
    except Exception:
        return 0


def pdf_page_text(raw: bytes, index: int) -> str:
    """Текстовый слой одной страницы PDF."""
    try:
        import pypdf

        rd = pypdf.PdfReader(io.BytesIO(raw))
        if index >= len(rd.pages):
            return ''
        return (rd.pages[index].extract_text() or '').strip()
    except Exception:
        return ''


def pdf_page_images(raw: bytes, index: int):
    """Картинки, вшитые в страницу PDF (сканы обычно одна картинка на лист)."""
    out = []
    try:
        import pypdf

        rd = pypdf.PdfReader(io.BytesIO(raw))
        if index >= len(rd.pages):
            return out
        for img in rd.pages[index].images:
            data = img.data
            if data and len(data) > 8000:
                out.append(data)
            if len(out) >= 3:
                break
    except Exception:
        return out
    return out


def read_page(raw: bytes, name: str, index: int, budget: int = 120):
    """Читает одну страницу файла: сначала текстовый слой, иначе распознаёт картинку.

    Возвращает (текст, способ): способ — 'text' или 'ocr'.
    """
    low = name.lower()

    if low.endswith(('.jpg', '.jpeg', '.png', '.webp', '.bmp', '.tif', '.tiff')):
        return ocr_image(prep_image(raw), 'image/jpeg', budget), 'ocr'

    if low.endswith('.pdf'):
        text = pdf_page_text(raw, index)
        if len(re.sub(r'\s', '', text)) >= 120:
            return text, 'text'
        chunks = []
        for img in pdf_page_images(raw, index):
            chunks.append(ocr_image(prep_image(img), 'image/jpeg', budget))
        joined = '\n'.join(c for c in chunks if c.strip())
        if joined.strip():
            return joined, 'ocr'
        return text, 'text'

    if low.endswith('.docx'):
        try:
            import docx

            d = docx.Document(io.BytesIO(raw))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(' | '.join(cells))
            return '\n'.join(parts), 'text'
        except Exception:
            return '', 'text'

    if low.endswith(('.txt', '.csv')):
        return raw.decode('utf-8', 'ignore'), 'text'

    return '', 'text'


def page_count(raw: bytes, name: str) -> int:
    low = name.lower()
    if low.endswith('.pdf'):
        return max(1, pdf_page_count(raw))
    return 1
