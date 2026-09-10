import base64
import io
import json
import os
import re
import ssl
import time
import uuid

import boto3
import psycopg2
import psycopg2.extras
import requests
import urllib3

import ocr

urllib3.disable_warnings()

CORS = {
    'Access-Control-Allow-Origin': '*',
    'Access-Control-Allow-Methods': 'GET, POST, DELETE, OPTIONS',
    'Access-Control-Allow-Headers': 'Content-Type, X-User-Id, X-Auth-Token, X-Session-Id',
    'Access-Control-Max-Age': '86400',
}

KIND_PD = 'project'
KIND_ID = 'executive'


def resp(code, data):
    return {
        'statusCode': code,
        'headers': {'Content-Type': 'application/json', **CORS},
        'body': json.dumps(data, ensure_ascii=False, default=str),
        'isBase64Encoded': False,
    }


def esc(v):
    return str(v).replace("'", "''")


def db():
    return psycopg2.connect(os.environ['DATABASE_URL'])


def rid(n=12):
    return uuid.uuid4().hex[:n]


def s3c():
    return boto3.client(
        's3',
        endpoint_url='https://bucket.poehali.dev',
        aws_access_key_id=os.environ['AWS_ACCESS_KEY_ID'],
        aws_secret_access_key=os.environ['AWS_SECRET_ACCESS_KEY'],
    )


def extract_text(name, raw):
    """Достаёт текст из PDF, DOCX или TXT. Возвращает (текст, число страниц)."""
    low = name.lower()
    try:
        if low.endswith('.pdf'):
            import pypdf

            rd = pypdf.PdfReader(io.BytesIO(raw))
            pages = len(rd.pages)
            out = []
            for p in rd.pages[:40]:
                out.append(p.extract_text() or '')
            return '\n'.join(out), pages
        if low.endswith('.docx'):
            import docx

            d = docx.Document(io.BytesIO(raw))
            paras = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        paras.append(' | '.join(cells))
            return '\n'.join(paras), max(1, len(paras) // 40)
        if low.endswith(('.txt', '.csv')):
            return raw.decode('utf-8', 'ignore'), 1
    except Exception as e:
        return f'[не удалось прочитать файл: {e}]', 0
    return '', 0


PROMPT_PD = """Ты — эксперт строительного контроля с 20-летним стажем. Проверяешь ПРОЕКТНУЮ документацию.

Объект: {obj}
Документы: {files}

ТЕКСТ ДОКУМЕНТАЦИИ:
{text}

Выполни ДВЕ проверки и верни СТРОГО JSON без пояснений.

ПРОВЕРКА 1 — соответствие нормам. Найди ошибки, нестыковки, противоречия, отсутствующие
обязательные разделы и решения. Каждое замечание обоснуй ссылкой на конкретный пункт
действующего норматива РФ (СП, ГОСТ, СНиП, ГрК РФ, ПП РФ 87 и т.п.) с цитатой требования.

ПРОВЕРКА 2 — контролепригодность. Оцени, может ли инспектор строительного контроля по
этой документации отследить и проконтролировать все этапы строительства: заданы ли
контролируемые параметры с допусками, указаны ли методы и средства контроля, объёмы
и моменты освидетельствования, перечень скрытых работ, требования к исполнительной
документации и геодезии.

Формат ответа:
{{"notes":[{{"severity":"Критическое|Замечание|Рекомендация","section":"раздел или лист",
"text":"суть нарушения","norm":"СП 48.13330.2019 п.5.3","quote":"цитата требования",
"demand":"что требуется сделать"}}],
"ctrl":[{{"severity":"Критическое|Замечание|Рекомендация","section":"раздел",
"text":"чего не хватает для контроля","norm":"пункт норматива","quote":"цитата",
"demand":"что требуется"}}],
"verdict":"вывод о соответствии нормам, 2-3 предложения",
"ctrl_verdict":"вывод о контролепригодности, 2-3 предложения",
"ctrl_score":число от 0 до 100}}

До 12 замечаний в notes и до 8 в ctrl. Только реальные нарушения, не выдумывай."""

PROMPT_ID = """Ты — эксперт строительного контроля с 20-летним стажем. Проверяешь ИСПОЛНИТЕЛЬНУЮ документацию.

Объект: {obj}
Документы: {files}

ТЕКСТ ДОКУМЕНТАЦИИ (распознан со сканов и фото, пометки: [от руки: ...], [печать: ...], [подпись]):
{text}

Учти: текст получен распознаванием сканов, поэтому отдельные символы могут быть искажены.
Не считай замечанием опечатку распознавания — оценивай смысл и содержание документа.
Отсутствие подписей, печатей и дат считай замечанием только если их явно нет в тексте.

Проверь по трём направлениям и верни СТРОГО JSON без пояснений:
1) соответствие требованиям проекта;
2) соответствие нормам и правилам (оформление актов освидетельствования скрытых работ,
   исполнительных схем, паспортов и сертификатов, журналов работ по СП 48.13330.2019,
   приказу Минстроя 1026/пр);
3) комплектность — все ли обязательные документы приложены.

Каждое замечание обоснуй ссылкой на конкретный пункт норматива РФ с цитатой.

Формат ответа:
{{"notes":[{{"severity":"Критическое|Замечание|Рекомендация","section":"документ или акт",
"text":"суть нарушения","norm":"СП 48.13330.2019 п.5.3","quote":"цитата требования",
"demand":"что требуется сделать"}}],
"complete":"вывод о комплектности: чего не хватает, 2-3 предложения",
"verdict":"вывод о качестве подготовки документации, 3-4 предложения",
"ctrl_score":число от 0 до 100}}

До 16 замечаний. Только реальные нарушения, не выдумывай."""


def parse_json(text):
    t = str(text).strip()
    t = re.sub(r'^```(?:json)?|```$', '', t, flags=re.M).strip()
    m = re.search(r'\{.*\}', t, re.S)
    if not m:
        raise ValueError('модель вернула не JSON')
    return json.loads(m.group(0))


def giga_token(auth):
    r = requests.post(
        'https://ngw.devices.sberbank.ru:9443/api/v2/oauth',
        data={'scope': 'GIGACHAT_API_PERS'},
        headers={
            'Authorization': f'Basic {auth}',
            'RqUID': str(uuid.uuid4()),
            'Accept': 'application/json',
        },
        timeout=10,
        verify=False,
    )
    r.raise_for_status()
    return r.json()['access_token']


def ask_cloudru(prompt, budget):
    key = os.environ.get('CLOUDRU_API_KEY', '')
    if not key:
        return None
    r = requests.post(
        'https://foundation-models.api.cloud.ru/v1/chat/completions',
        json={
            'model': 'GigaChat/GigaChat-2-Max',
            'temperature': 0.1,
            'max_tokens': 6000,
            'messages': [{'role': 'user', 'content': prompt}],
        },
        headers={'Authorization': f'Bearer {key}'},
        timeout=budget,
    )
    r.raise_for_status()
    return parse_json(r.json()['choices'][0]['message']['content'])


def ask_deepseek(prompt, budget):
    key = os.environ.get('DEEPSEEK_API_KEY', '')
    if not key:
        return None
    r = requests.post(
        'https://api.deepseek.com/chat/completions',
        json={
            'model': 'deepseek-chat',
            'temperature': 0.1,
            'max_tokens': 6000,
            'messages': [{'role': 'user', 'content': prompt}],
        },
        headers={'Authorization': f'Bearer {key}'},
        timeout=budget,
    )
    r.raise_for_status()
    return parse_json(r.json()['choices'][0]['message']['content'])


def ask_gigachat(prompt, budget):
    auth = os.environ.get('GIGACHAT_AUTH_KEY', '')
    if not auth:
        return None
    started = time.time()
    token = giga_token(auth)
    left = budget - (time.time() - started)
    if left < 2:
        raise TimeoutError('не хватило времени на запрос к ИИ')
    r = requests.post(
        'https://gigachat.devices.sberbank.ru/api/v1/chat/completions',
        json={
            'model': 'GigaChat-2',
            'temperature': 0.1,
            'max_tokens': 4000,
            'messages': [{'role': 'user', 'content': prompt}],
        },
        headers={'Authorization': f'Bearer {token}', 'Accept': 'application/json'},
        timeout=left,
        verify=False,
    )
    r.raise_for_status()
    return parse_json(r.json()['choices'][0]['message']['content'])


PROVIDERS = (('cloudru', ask_cloudru), ('deepseek', ask_deepseek), ('gigachat', ask_gigachat))


def run_ai(prompt, budget):
    errors = []
    for name, fn in PROVIDERS:
        try:
            out = fn(prompt, budget)
            if out:
                return out, name, ''
        except Exception as e:
            errors.append(f'{name}: {e}')
    return None, '', '; '.join(errors) or 'не подключён ни один ИИ-ключ'


def save_notes(cur, review_id, items, scope):
    n = 0
    for it in items or []:
        n += 1
        cur.execute(
            'INSERT INTO doc_review_notes (id, review_id, scope, num, severity, section, '
            'text, norm_ref, norm_quote, demand) VALUES ('
            f"'{esc(rid())}', '{esc(review_id)}', '{esc(scope)}', {n}, "
            f"'{esc(str(it.get('severity', 'Замечание'))[:12])}', "
            f"'{esc(str(it.get('section', ''))[:300])}', "
            f"'{esc(str(it.get('text', ''))[:2000])}', "
            f"'{esc(str(it.get('norm', ''))[:300])}', "
            f"'{esc(str(it.get('quote', ''))[:1500])}', "
            f"'{esc(str(it.get('demand', ''))[:1500])}')"
        )
    return n


def to_review(r):
    return {
        'id': r['id'],
        'kind': r['kind'],
        'objectId': r['object_id'] or '',
        'objectName': r['object_name'],
        'title': r['title'],
        'inspector': r['inspector'],
        'status': r['status'],
        'filesCount': r['files_count'],
        'pagesCount': r['pages_count'],
        'verdict': r['verdict'],
        'ctrlVerdict': r['ctrl_verdict'],
        'ctrlScore': r['ctrl_score'],
        'completeNote': r['complete_note'],
        'ocrPages': r.get('ocr_pages', 0),
        'engine': r['engine'],
        'error': r['error'],
        'createdAt': r['created_at'],
        'checkedAt': r['checked_at'],
    }


def handler(event: dict, context) -> dict:
    """Входной контроль проектной и исполнительной документации силами ИИ.

    Загружает файлы в хранилище, извлекает текст, проверяет на соответствие
    нормам, контролепригодность и комплектность, формирует акты с замечаниями
    и ссылками на пункты нормативных документов.
    """
    method = event.get('httpMethod', 'GET')
    if method == 'OPTIONS':
        return {'statusCode': 200, 'headers': CORS, 'body': ''}

    params = event.get('queryStringParameters') or {}
    body = {}
    if event.get('body'):
        try:
            body = json.loads(event['body'])
        except Exception:
            body = {}
    action = body.get('action') or params.get('action') or ''

    if action == 'visiontest':
        import base64 as _b
        from PIL import Image, ImageDraw
        im = Image.new('RGB', (600, 200), (255, 255, 255))
        ImageDraw.Draw(im).text((40, 80), body.get('word', 'ARBUZ-7315'), fill=(0, 0, 0))
        buf = io.BytesIO(); im.save(buf, 'PNG')
        b64 = _b.b64encode(buf.getvalue()).decode()
        out = {}
        for m in (body.get('models') or ['qwen/qwen3-vl-30b-a3b-instruct']):
            try:
                r = requests.post(
                    'https://foundation-models.api.cloud.ru/v1/chat/completions',
                    json={'model': m, 'max_tokens': 40, 'temperature': 0, 'messages': [{
                        'role': 'user', 'content': [
                            {'type': 'text', 'text': 'Какой текст на картинке? Ответь только текстом.'},
                            {'type': 'image_url', 'image_url': {'url': f'data:image/png;base64,{b64}'}}]}]},
                    headers={'Authorization': f"Bearer {os.environ.get('CLOUDRU_API_KEY','')}"},
                    timeout=60,
                )
                j = r.json()
                out[m] = j['choices'][0]['message']['content'][:90] if r.status_code == 200 else f'{r.status_code} {r.text[:90]}'
            except Exception as e:
                out[m] = str(e)[:90]
        return resp(200, out)

    if action == 'ocrtest':
        import ocr as _o
        c2 = db(); k2 = c2.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        k2.execute(
            "SELECT url FROM doc_review_files WHERE review_id = '"
            + esc(body.get('reviewId', '')) + "' ORDER BY created_at LIMIT 1"
        )
        row = k2.fetchone(); k2.close(); c2.close()
        if not row:
            return resp(200, {'error': 'нет файлов'})
        raw = s3c().get_object(Bucket='files', Key=row['url'].split('/bucket/', 1)[-1])['Body'].read()
        started = time.time()
        try:
            txt = _o.ocr_image(_o.prep_image(raw), 'image/jpeg', budget=int(body.get('budget', 25)))
            return resp(200, {'sec': round(time.time() - started, 1), 'chars': len(txt), 'text': txt[:1200]})
        except Exception as e:
            return resp(200, {'sec': round(time.time() - started, 1), 'err': str(e)[:300]})

    if action == 'ping2':
        key = os.environ.get('CLOUDRU_API_KEY', '')
        out = {'keyLen': len(key), 'keyHead': key[:6] if key else ''}
        try:
            r = requests.get(
                'https://foundation-models.api.cloud.ru/v1/models',
                headers={'Authorization': f'Bearer {key}'},
                timeout=3,
            )
            out['models'] = r.status_code
        except Exception as e:
            out['models'] = str(e)[:120]
        for m in (body.get('models') or ['openai/gpt-4o-mini']):
            try:
                r = requests.post(
                    'https://foundation-models.api.cloud.ru/v1/chat/completions',
                    json={
                        'model': m,
                        'max_tokens': 5,
                        'messages': [{'role': 'user', 'content': 'hi'}],
                    },
                    headers={'Authorization': f'Bearer {key}'},
                    timeout=3.5,
                )
                out[m] = f'{r.status_code} {r.text[:110]}'
            except Exception as e:
                out[m] = str(e)[:110]
        return resp(200, out)

    if action == 'ping':
        key = os.environ.get('CLOUDRU_API_KEY', '')
        if not key:
            return resp(200, {'ok': False, 'message': 'Ключ Cloud.ru не задан'})
        started = time.time()
        try:
            r = requests.post(
                'https://foundation-models.api.cloud.ru/v1/chat/completions',
                json={
                    'model': 'GigaChat/GigaChat-2-Max',
                    'max_tokens': 5,
                    'messages': [{'role': 'user', 'content': 'привет'}],
                },
                headers={'Authorization': f'Bearer {key}'},
                timeout=3.5,
            )
            return resp(
                200,
                {
                    'ok': r.status_code == 200,
                    'http': r.status_code,
                    'sec': round(time.time() - started, 1),
                    'body': r.text[:400],
                },
            )
        except Exception as e:
            return resp(200, {'ok': False, 'message': f'{type(e).__name__}: {str(e)[:250]}'})

    if action == 'providers':
        out = {'ocr': 'ключ Cloud.ru задан' if os.environ.get('CLOUDRU_API_KEY') else 'ключ Cloud.ru НЕ задан'}
        for name, fn in PROVIDERS:
            started = time.time()
            try:
                r = fn('Ответь строго JSON: {"ok":1}', 20)
                out[name] = 'ключ не задан' if r is None else f'работает за {time.time() - started:.1f}с'
            except Exception as e:
                out[name] = f'{type(e).__name__}: {str(e)[:160]}'
        return resp(200, out)

    conn = db()
    cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    try:
        if method == 'GET' and action == 'one':
            rv = params.get('id', '')
            cur.execute(f"SELECT * FROM doc_reviews WHERE id = '{esc(rv)}'")
            row = cur.fetchone()
            if not row:
                return resp(404, {'error': 'not_found'})
            cur.execute(
                f"SELECT * FROM doc_review_notes WHERE review_id = '{esc(rv)}' "
                'ORDER BY scope, num'
            )
            notes = [
                {
                    'id': n['id'],
                    'scope': n['scope'],
                    'num': n['num'],
                    'severity': n['severity'],
                    'section': n['section'],
                    'text': n['text'],
                    'norm': n['norm_ref'],
                    'quote': n['norm_quote'],
                    'demand': n['demand'],
                }
                for n in cur.fetchall()
            ]
            cur.execute(
                f"SELECT * FROM doc_review_files WHERE review_id = '{esc(rv)}' "
                'ORDER BY created_at'
            )
            files = [
                {
                    'id': f['id'],
                    'name': f['name'],
                    'url': f['url'],
                    'sizeKb': f['size_kb'],
                    'pages': f['pages'],
                }
                for f in cur.fetchall()
            ]
            return resp(200, {'item': to_review(row), 'notes': notes, 'files': files})

        if method == 'GET':
            kind = params.get('kind', KIND_PD)
            cur.execute(
                f"SELECT * FROM doc_reviews WHERE kind = '{esc(kind)}' "
                'ORDER BY created_at DESC LIMIT 300'
            )
            return resp(200, {'items': [to_review(r) for r in cur.fetchall()]})

        if method == 'POST' and action == 'create':
            kind = body.get('kind', KIND_PD)
            new_id = rid()
            cur.execute(
                'INSERT INTO doc_reviews (id, kind, object_id, object_name, title, inspector) '
                f"VALUES ('{esc(new_id)}', '{esc(kind)}', "
                f"'{esc(body.get('objectId', ''))}', '{esc(body.get('objectName', ''))}', "
                f"'{esc(body.get('title', ''))}', '{esc(body.get('inspector', ''))}')"
            )
            conn.commit()
            return resp(200, {'id': new_id})

        if method == 'POST' and action == 'upload':
            rv = body.get('reviewId', '')
            name = str(body.get('name', 'файл'))[:300]
            content = body.get('content', '')
            if not rv or not content:
                return resp(400, {'error': 'review_and_content_required'})
            raw = base64.b64decode(str(content).split(',')[-1])
            fid = rid()
            safe = re.sub(r'[^\w.\-]+', '_', name)[:80]
            key = f'docaudit/{rv}/{fid}_{safe}'
            s3c().put_object(
                Bucket='files',
                Key=key,
                Body=raw,
                ContentType=body.get('mime') or 'application/octet-stream',
            )
            url = f"https://cdn.poehali.dev/projects/{os.environ['AWS_ACCESS_KEY_ID']}/bucket/{key}"
            pages = ocr.page_count(raw, name)
            cur.execute(
                'INSERT INTO doc_review_files (id, review_id, name, url, mime, size_kb, pages) '
                f"VALUES ('{esc(fid)}', '{esc(rv)}', '{esc(name)}', '{esc(url)}', "
                f"'{esc(body.get('mime', ''))}', {len(raw) // 1024}, {pages})"
            )
            cur.execute(
                f"UPDATE doc_reviews SET files_count = files_count + 1, "
                f"pages_count = pages_count + {pages}, stage = 'read' "
                f"WHERE id = '{esc(rv)}'"
            )
            conn.commit()
            return resp(
                200,
                {'id': fid, 'name': name, 'url': url, 'pages': pages, 'sizeKb': len(raw) // 1024},
            )

        if method == 'POST' and action == 'page':
            rv = body.get('reviewId', '')
            cur.execute(
                'SELECT f.* FROM doc_review_files f '
                f"WHERE f.review_id = '{esc(rv)}' AND f.done_pages < f.pages "
                'ORDER BY f.created_at LIMIT 1'
            )
            f = cur.fetchone()
            if not f:
                cur.execute(
                    f"SELECT COALESCE(SUM(LENGTH(text)), 0) AS chars FROM doc_review_pages "
                    f"WHERE review_id = '{esc(rv)}'"
                )
                chars = cur.fetchone()['chars']
                cur.execute(
                    f"UPDATE doc_reviews SET stage = 'ready' WHERE id = '{esc(rv)}'"
                )
                conn.commit()
                return resp(200, {'done': True, 'chars': chars})

            idx = f['done_pages']
            s3key = f['url'].split('/bucket/', 1)[-1]
            raw = s3c().get_object(Bucket='files', Key=s3key)['Body'].read()
            try:
                text, how = ocr.read_page(raw, f['name'], idx, budget=110)
            except Exception as e:
                text, how = '', 'text'
                cur.execute(
                    f"UPDATE doc_reviews SET error = '{esc(str(e)[:500])}' WHERE id = '{esc(rv)}'"
                )

            cur.execute(
                'INSERT INTO doc_review_pages (id, review_id, file_id, page_no, method, text) '
                f"VALUES ('{esc(rid(18))}', '{esc(rv)}', '{esc(f['id'])}', {idx + 1}, "
                f"'{esc(how)}', '{esc(text[:30000])}')"
            )
            cur.execute(
                f"UPDATE doc_review_files SET done_pages = done_pages + 1, "
                f"ocr_pages = ocr_pages + {1 if how == 'ocr' else 0} "
                f"WHERE id = '{esc(f['id'])}'"
            )
            cur.execute(
                f"UPDATE doc_reviews SET done_pages = done_pages + 1, "
                f"ocr_pages = ocr_pages + {1 if how == 'ocr' else 0} "
                f"WHERE id = '{esc(rv)}'"
            )
            conn.commit()
            cur.execute(
                f"SELECT done_pages, pages_count FROM doc_reviews WHERE id = '{esc(rv)}'"
            )
            st = cur.fetchone()
            return resp(
                200,
                {
                    'done': False,
                    'page': idx + 1,
                    'file': f['name'],
                    'method': how,
                    'chars': len(text),
                    'donePages': st['done_pages'],
                    'totalPages': st['pages_count'],
                },
            )

        if method == 'POST' and action == 'analyze':
            rv = body.get('reviewId', '')
            cur.execute(f"SELECT * FROM doc_reviews WHERE id = '{esc(rv)}'")
            review = cur.fetchone()
            if not review:
                return resp(404, {'error': 'not_found'})

            cur.execute(
                f"SELECT name FROM doc_review_files WHERE review_id = '{esc(rv)}' "
                'ORDER BY created_at'
            )
            files = cur.fetchall()
            if not files:
                return resp(400, {'error': 'no_files'})

            cur.execute(
                'SELECT p.page_no, p.text, f.name FROM doc_review_pages p '
                'JOIN doc_review_files f ON f.id = p.file_id '
                f"WHERE p.review_id = '{esc(rv)}' AND LENGTH(p.text) > 40 "
                'ORDER BY f.created_at, p.page_no'
            )
            rows = cur.fetchall()
            chunks = [f"=== {r['name']}, лист {r['page_no']} ===\n{r['text'][:9000]}" for r in rows]

            joined = '\n\n'.join(chunks)[:60000]
            if not joined.strip():
                cur.execute(
                    "UPDATE doc_reviews SET status = 'error', error = "
                    "'Не удалось прочитать документацию' "
                    f"WHERE id = '{esc(rv)}'"
                )
                conn.commit()
                return resp(200, {'error': 'no_text', 'message': 'Не удалось прочитать текст'})

            names = ', '.join(f['name'] for f in files)
            tpl = PROMPT_PD if review['kind'] == KIND_PD else PROMPT_ID
            prompt = tpl.format(obj=review['object_name'] or '—', files=names, text=joined)

            data, engine, err = run_ai(prompt, 170)
            if not data:
                cur.execute(
                    f"UPDATE doc_reviews SET status = 'error', error = '{esc(err[:900])}' "
                    f"WHERE id = '{esc(rv)}'"
                )
                conn.commit()
                return resp(200, {'error': 'ai_failed', 'message': err})

            cur.execute(f"DELETE FROM doc_review_notes WHERE review_id = '{esc(rv)}'")
            n1 = save_notes(cur, rv, data.get('notes'), 'norms')
            n2 = save_notes(cur, rv, data.get('ctrl'), 'ctrl')
            score = data.get('ctrl_score', 0)
            try:
                score = max(0, min(100, int(score)))
            except Exception:
                score = 0
            cur.execute(
                "UPDATE doc_reviews SET status = 'done', error = '', "
                f"verdict = '{esc(str(data.get('verdict', ''))[:3000])}', "
                f"ctrl_verdict = '{esc(str(data.get('ctrl_verdict', ''))[:3000])}', "
                f"complete_note = '{esc(str(data.get('complete', ''))[:3000])}', "
                f"ctrl_score = {score}, engine = '{esc(engine)}', checked_at = now() "
                f"WHERE id = '{esc(rv)}'"
            )
            conn.commit()
            return resp(200, {'ok': True, 'notes': n1, 'ctrlNotes': n2, 'engine': engine})

        if method == 'DELETE':
            rv = params.get('id', '') or body.get('id', '')
            cur.execute(f"DELETE FROM doc_review_notes WHERE review_id = '{esc(rv)}'")
            cur.execute(f"DELETE FROM doc_review_files WHERE review_id = '{esc(rv)}'")
            cur.execute(f"DELETE FROM doc_reviews WHERE id = '{esc(rv)}'")
            conn.commit()
            return resp(200, {'ok': True})

        return resp(400, {'error': 'unknown_action'})
    finally:
        cur.close()
        conn.close()