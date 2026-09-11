"""Сборка акта проверки документации в формате Word (.docx)."""

import io

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x3A, 0x6E)
GREY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xB3, 0x26, 0x1E)

MONTHS = [
    'января', 'февраля', 'марта', 'апреля', 'мая', 'июня',
    'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря',
]


def ru_date(value):
    if not value:
        return '—'
    try:
        return f'{value.day} {MONTHS[value.month - 1]} {value.year} г.'
    except Exception:
        return str(value)[:10]


def shade(cell, color):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(el)


def set_font(doc):
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    fmt = style.paragraph_format
    fmt.space_after = Pt(4)
    fmt.line_spacing = 1.15


def para(doc, text='', size=12, bold=False, italic=False, color=None,
         align=None, space_before=0, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return p


def kv_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for name, value in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(5.5)
        cells[1].width = Cm(11.5)
        shade(cells[0], 'F2F4F8')
        for cell, text, bold in ((cells[0], name, False), (cells[1], value or '—', True)):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.bold = bold
            run.font.size = Pt(11)
    return table


def notes_block(doc, notes):
    if not notes:
        para(doc, 'Замечаний не выявлено.', italic=True, color=GREY, space_before=4)
        return

    for note in notes:
        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(10)
        head.paragraph_format.space_after = Pt(2)
        head.paragraph_format.keep_with_next = True

        num = head.add_run(f"Замечание № {note.get('num', '')}. ")
        num.bold = True
        num.font.size = Pt(12)
        num.font.color.rgb = ACCENT

        severity = str(note.get('severity') or 'Замечание')
        sev = head.add_run(f'[{severity}]')
        sev.bold = True
        sev.font.size = Pt(10)
        sev.font.color.rgb = RED if severity.startswith('Крит') else GREY

        section = str(note.get('section') or '')
        if section:
            src = head.add_run(f'  ({section})')
            src.font.size = Pt(10)
            src.font.color.rgb = GREY

        body = doc.add_paragraph()
        body.paragraph_format.space_after = Pt(3)
        body.paragraph_format.left_indent = Cm(0.6)
        body.add_run(str(note.get('text') or '')).font.size = Pt(12)

        norm = str(note.get('norm') or '')
        quote = str(note.get('quote') or '')
        if norm or quote:
            ref = doc.add_paragraph()
            ref.paragraph_format.left_indent = Cm(0.6)
            ref.paragraph_format.space_after = Pt(3)
            if norm:
                r = ref.add_run(f'Основание: {norm}')
                r.bold = True
                r.font.size = Pt(10.5)
            if quote:
                q = ref.add_run(f'\n«{quote}»')
                q.italic = True
                q.font.size = Pt(10.5)
                q.font.color.rgb = GREY

        demand = str(note.get('demand') or '')
        if demand:
            req = doc.add_paragraph()
            req.paragraph_format.left_indent = Cm(0.6)
            req.paragraph_format.space_after = Pt(3)
            lbl = req.add_run('Требуется: ')
            lbl.bold = True
            lbl.font.size = Pt(11)
            req.add_run(demand).font.size = Pt(11)


def act_section(doc, title, subtitle, notes, verdict, score=None,
                score_label='', extra_title='', extra_text=''):
    para(doc, title, size=13, bold=True, color=ACCENT,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=2)
    if subtitle:
        para(doc, subtitle, size=10.5, italic=True, color=GREY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    if score is not None:
        para(doc, f'{score_label}: {score} %', size=12, bold=True, space_after=6)

    para(doc, f'Выявлено замечаний: {len(notes)}', size=11, bold=True, space_before=4)
    notes_block(doc, notes)

    if extra_text:
        para(doc, extra_title, size=11, bold=True, space_before=12, space_after=2)
        para(doc, extra_text, size=11.5)

    para(doc, 'Вывод', size=11, bold=True, space_before=12, space_after=2)
    para(doc, verdict or 'Вывод не сформирован.', size=11.5)


def signatures(doc, inspector):
    para(doc, 'Проверку провёл:', size=11, bold=True, space_before=20, space_after=8)

    table = doc.add_table(rows=1, cols=2)
    cells = table.rows[0].cells
    cells[0].width = Cm(9)
    cells[1].width = Cm(8)

    left = cells[0].paragraphs[0]
    left.add_run(inspector or '________________________').font.size = Pt(11)
    hint = cells[0].add_paragraph()
    h = hint.add_run('(должность, фамилия, инициалы)')
    h.font.size = Pt(9)
    h.font.color.rgb = GREY

    right = cells[1].paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.add_run('_____________________').font.size = Pt(11)
    sign = cells[1].add_paragraph()
    sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    s = sign.add_run('(подпись)')
    s.font.size = Pt(9)
    s.font.color.rgb = GREY


def build_docx(review, notes, files, is_pd):
    """Формирует .docx акта проверки и возвращает его байты."""
    doc = Document()
    set_font(doc)

    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    para(doc, 'АКТ ПРОВЕРКИ ДОКУМЕНТАЦИИ', size=16, bold=True, color=ACCENT,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc,
         'проектной документации' if is_pd else 'исполнительной документации',
         size=11.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

    norm_notes = [n for n in notes if n.get('scope') == 'norms']
    ctrl_notes = [n for n in notes if n.get('scope') == 'ctrl']

    kv_table(doc, [
        ('Объект', review.get('objectName')),
        ('Предмет проверки', review.get('title')),
        ('Проверку провёл', review.get('inspector')),
        ('Дата проверки', ru_date(review.get('checkedAt') or review.get('createdAt'))),
        ('Проверено файлов', str(review.get('filesCount') or 0)),
        ('Проверено листов', str(review.get('pagesCount') or 0)),
        ('Всего замечаний', str(len(notes))),
    ])

    if files:
        para(doc, 'Проверенные документы', size=11, bold=True,
             space_before=14, space_after=4)
        for i, f in enumerate(files, 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            pages = f.get('pages') or 0
            tail = f', листов: {pages}' if pages else ''
            p.add_run(f"{i}. {f.get('name', '')}{tail}").font.size = Pt(11)

    act_section(
        doc,
        'Соответствие требованиям норм и правил',
        'Проверка решений документации на соответствие действующей нормативной базе',
        norm_notes,
        review.get('verdict', ''),
        extra_title='Комплектность',
        extra_text='' if is_pd else str(review.get('completeNote') or ''),
    )

    if is_pd:
        act_section(
            doc,
            'Оценка контролепригодности',
            'Возможность контроля всех этапов строительства по данной документации',
            ctrl_notes,
            review.get('ctrlVerdict', ''),
            score=review.get('ctrlScore', 0),
            score_label='Оценка контролепригодности',
        )

    signatures(doc, review.get('inspector', ''))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
