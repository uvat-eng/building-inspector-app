import datetime
import io
import re

import openpyxl

CATEGORY_COLS = {
    28: 'ОТ и ТБ',
    30: 'Аттестация персонала',
    32: 'Разрешительная документация',
    34: 'Исполнительная документация',
    36: 'Складирование и транспортировка',
    38: 'Входной контроль',
    40: 'Общестроительные работы',
    42: 'Сборка и сварка',
    44: 'Электромонтажные работы',
    46: 'Оборудование и инструмент',
    48: 'Отступления от проектных решений / Согласование',
}

NATURE_COLS = {25: 'ОТ и ПБ', 26: 'Технология и качество', 27: 'Документация'}

COL = {
    'num': 1,
    'contractor': 2,
    'point': 3,
    'content': 4,
    'normRef': 5,
    'docRef': 6,
    'place': 7,
    'orderNo': 8,
    'issuedAt': 9,
    'dueAt': 10,
    'factAt': 11,
    'extension': 12,
    'status': 13,
    'responsible': 14,
    'inspector': 15,
    'stopWork': 17,
}


def txt(v) -> str:
    if v is None:
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%d.%m.%Y')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return '' if s.startswith('=') else s


def iso(v) -> str:
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    s = txt(v)
    m = re.match(r'^(\d{2})\.(\d{2})\.(\d{4})$', s)
    if m:
        return f'{m.group(3)}-{m.group(2)}-{m.group(1)}'
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
    return m.group(0) if m else ''


def num(v) -> float:
    try:
        f = float(v)
        return f
    except (TypeError, ValueError):
        return 0.0


def find_header(ws) -> int:
    """Ищет строку заголовка таблицы предписаний."""
    for r in range(1, min(ws.max_row, 60) + 1):
        joined = ' '.join(txt(ws.cell(r, c).value).lower() for c in range(1, 10))
        if 'наименование подрядной организации' in joined or (
            'содержание замечан' in joined and 'подрядн' in joined
        ):
            return r
    return 17


def data_start(ws, header_row: int) -> int:
    """Строка, с которой начинаются данные (пропускает нумерацию колонок)."""
    for r in range(header_row + 1, min(header_row + 8, ws.max_row) + 1):
        a = txt(ws.cell(r, COL['num']).value)
        b = txt(ws.cell(r, COL['contractor']).value)
        if b and not b.isdigit():
            return r
        if a == '1' and txt(ws.cell(r, 2).value) == '2':
            return r + 1
    return header_row + 4


def parse_workbook(raw: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    ws = wb.worksheets[0]

    merged = {}
    for rng in ws.merged_cells.ranges:
        top = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = top

    def cell(r, c):
        v = ws.cell(r, c).value
        if v in (None, '') and (r, c) in merged:
            v = merged[(r, c)]
        return v

    header = find_header(ws)
    start = data_start(ws, header)

    rows = []
    contractor = ''
    report_date = ''
    author = ''

    for r in range(start, ws.max_row + 1):
        content = txt(cell(r, COL['content']))
        c_name = txt(cell(r, COL['contractor']))
        if c_name:
            contractor = c_name
        if not content and not txt(cell(r, COL['orderNo'])):
            continue

        nature = ''
        for c, name in NATURE_COLS.items():
            if num(ws.cell(r, c).value) > 0:
                nature = name
                break

        category = ''
        for c, name in CATEGORY_COLS.items():
            if num(ws.cell(r, c).value) > 0:
                category = name
                break

        status_raw = txt(cell(r, COL['status'])).lower()
        status = 'Устранено' if status_raw.startswith('устранено') else 'Не устранено'
        inspector = txt(cell(r, COL['inspector']))
        if inspector and not author:
            author = inspector

        issued = iso(cell(r, COL['issuedAt']))
        if issued and (not report_date or issued > report_date):
            report_date = issued

        rows.append(
            {
                'id': f'x{r}',
                'contractor': contractor,
                'point': int(num(cell(r, COL['point'])) or 1),
                'content': content,
                'normRef': txt(cell(r, COL['normRef'])),
                'docRef': txt(cell(r, COL['docRef'])),
                'place': txt(cell(r, COL['place'])),
                'orderNo': txt(cell(r, COL['orderNo'])),
                'issuedAt': issued,
                'dueAt': iso(cell(r, COL['dueAt'])),
                'factAt': iso(cell(r, COL['factAt'])),
                'extension': txt(cell(r, COL['extension'])),
                'status': status,
                'responsible': txt(cell(r, COL['responsible'])),
                'inspector': inspector,
                'stopWork': num(ws.cell(r, COL['stopWork']).value) > 0,
                'nature': nature,
                'category': category,
                'photos': [],
            }
        )

    return {'rows': rows, 'date': report_date, 'author': author, 'sheet': ws.title}
