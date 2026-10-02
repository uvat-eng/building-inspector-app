"""Сборка предписания в формате Word (.docx).

Документ повторяет утверждённую форму предписания: тот же порядок блоков,
те же формулировки, та же таблица нарушений и те же подписи, что в бланке.
Собирается настоящий Word — HTML с расширением .doc телефон считает битым.
"""

import io

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import tenant

MONTHS = [
    'января', 'февраля', 'марта', 'апреля', 'мая', 'июня',
    'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря',
]



def ru_long(value):
    """Дата словами: 30 сентября 2026 г."""
    if not value:
        return ''
    text = str(value)[:10]
    try:
        if '.' in text:
            day, month, year = text.split('.')
        else:
            year, month, day = text.split('-')
        return f'{int(day)} {MONTHS[int(month) - 1]} {year} г.'
    except Exception:
        return text


def line(value, min_len=24):
    """Значение или пустая линия для заполнения от руки — как в бланке."""
    text = str(value or '').strip()
    return text if text else '_' * min_len


def set_font(doc):
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    fmt = style.paragraph_format
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.15


def para(doc, text='', size=11, bold=False, italic=False, align=None,
         space_before=0, space_after=3):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
    return p


def hint(doc, text):
    """Пояснение под строкой мелким курсивом — как в бланке."""
    return para(doc, text, size=8.5, italic=True,
                align=WD_ALIGN_PARAGRAPH.CENTER, space_after=5)


def set_col_widths(table, widths):
    """Ширины столбцов: Word держит их только при явной установке в каждой ячейке."""
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def violations_table(doc, items, order):
    """Таблица нарушений — пять граф, как в утверждённой форме."""
    titles = (
        '№ п/п',
        'Краткое изложение выявленного нарушения с указанием места обнаружения',
        'Пункт требований тех. регламентов, иных нормативных правовых актов, '
        'проектной и рабочей документации, требования которых не исполнены',
        'Предлагаемые меры и срок устранения нарушения',
        '№ фотодокумента',
    )
    widths = (Cm(1.1), Cm(6.0), Cm(5.2), Cm(3.6), Cm(1.9))

    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    for cell, title in zip(table.rows[0].cells, titles):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(8)

    if items:
        for i, it in enumerate(items, 1):
            cells = table.add_row().cells
            deadline = it.get('deadline') or order.get('deadline') or '—'
            values = (
                str(i),
                str(it.get('title') or ''),
                str(it.get('normRef') or '—'),
                f'Устранить нарушение.\nСрок: {deadline}',
                f'фото {i}' if it.get('photos') else '—',
            )
            for j, (cell, text) in enumerate(zip(cells, values)):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                if j in (0, 4):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text)
                run.font.size = Pt(9)
    else:
        cells = table.add_row().cells
        merged = cells[0].merge(cells[4])
        p = merged.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run('нарушений не выявлено').font.size = Pt(9)

    set_col_widths(table, widths)
    return table


def sign_table(doc, rows):
    """Блок подписей — две графы, как в бланке."""
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for left, right, small in rows:
        cells = table.add_row().cells
        for cell, text in ((cells[0], left), (cells[1], right)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(text))
            run.font.size = Pt(8.5) if small else Pt(10)
            run.italic = small
    set_col_widths(table, (Cm(9.8), Cm(8.0)))
    return table


def build_order_docx(order, contractor=None):
    """Формирует .docx предписания по утверждённой форме и возвращает байты."""
    doc = Document()
    set_font(doc)

    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

    body = order.get('body') or {}
    items = body.get('items') or []
    contractor = contractor or {}
    org = body.get('inspectionOrg') or tenant.org_name()
    issued_to = order.get('issuedTo') or contractor.get('name') or ''
    number = order.get('number') or ''

    # Шапка
    para(doc, org, bold=True)
    para(doc, f"Объект строительства: «{body.get('objectTitle') or ''}»", bold=True)
    para(doc, f"Шифр объекта: {line(body.get('objectCode'), 20)}", bold=True)
    para(doc,
         f"{ru_long(order.get('createdAt'))}        Время {body.get('time') or '__:__'}",
         space_before=8)

    para(doc, f'ПРЕДПИСАНИЕ № {number}', size=13, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=8)

    # Кому выдано
    para(doc, f'Выдано {line(issued_to, 46)}')
    hint(doc, '(наименование организации, осуществляющей строительство)')
    para(doc, 'осуществляющей строительство на основании договора подряда')
    para(doc, f"№ {line(body.get('contractNo'), 22)} от {line(body.get('contractDate'), 14)}")
    hint(doc, '(№ договора подряда и дата заключения договора)')

    # Кем выдано
    para(doc,
         'Мною, представителем независимого строительного контроля: '
         f"Инспектор СК {org} {line(order.get('inspector'), 26)}",
         space_before=6)
    hint(doc, '(фамилия И.О. ответственного представителя организации, '
              'осуществляющей строительный контроль)')
    para(doc,
         f"На основании документа № {line(body.get('assignDocNo'), 10)} "
         f"от {line(body.get('assignDocDate'), 14)}")
    hint(doc, '(№ и дата распорядительного документа о назначении специалиста НСК на объект)')

    # В присутствии
    para(doc, 'В присутствии:', space_before=6)
    para(doc, 'Ответственного представителя застройщика/технического заказчика:')
    para(doc, line(body.get('customerRep'), 60))
    hint(doc, '(должность, Фамилия И.О.)')
    para(doc, 'Ответственного представителя строительного подрядчика, '
              'осуществляющего строительство:')
    para(doc, line(body.get('contractorRep'), 60))
    hint(doc, '(наименование лица, осуществляющего строительство (генеральный строительный '
              'подрядчик), должность, Ф.И.О представителя)')

    para(doc,
         'По результатам проведенной проверки соответствия выполняемых работ требованиям '
         'нормативных актов, ПД, РД, требованиям технических регламентов, требованиям '
         'градостроительного плана земельного участка выявлены следующие нарушения:',
         space_before=8, space_after=6)

    violations_table(doc, items, order)

    # Предлагаемые меры
    count = len(items)
    if count == 0:
        rng = '—'
    elif count == 1:
        rng = '№1'
    else:
        rng = f'№1 — №{count}'

    para(doc, 'Предлагаю:', space_before=8)
    para(doc, f"Устранить несоответствия {rng} до {order.get('deadline') or '__.__.____'}.")
    para(doc,
         'Срок исполнения (в случае остановки работ — до устранения нарушения) '
         f"{line(body.get('stopNote'), 12)}.")
    para(doc,
         'Остановить производство (заполнить в случае остановки работ) '
         f"{line(body.get('stopWorks'), 14)}.")
    para(doc,
         'Вид работ, отдельный этап работ (земляные, сварочные, АКЗ и т. д.): '
         f"{line(body.get('workType'), 26)}")
    para(doc, f"Приложение: фотоматериалы на {count or '__'} л.")

    # Правовые основания
    para(doc,
         'В связи с тем, что выявленные в ходе проверки факты повлекли нарушения нормативных '
         'и правовых актов, руководствуясь статьей 53 Градостроительного кодекса Российской '
         'Федерации от 29.12.2004 № 190-ФЗ и статьями 705, 706, 714, 715, 720, 721, 723, 745, '
         '748, 751, 753, 754, 755 Гражданского кодекса Российской Федерации данное предписание '
         'может служить основанием для остановки работ и ведения претензионной работы.',
         size=9, space_before=10)
    para(doc,
         'Вам предписывается устранить вышеуказанные нарушения в установленные для этого сроки '
         'и направить Акт об устранении каждого пункта настоящего предписания с перечислением '
         'принятых мер, подтверждающих факт устранения нарушений, ответственному представителю '
         'организации по независимому строительному контролю для освидетельствования устранения '
         'выявленных нарушений по настоящему предписанию.',
         size=9, space_before=6)

    # Принял к исполнению
    para(doc, f'Предписание № {number} к исполнению принял:', space_before=10)
    para(doc, 'Ответственный представитель лица, осуществляющего строительство')
    para(doc, line(body.get('contractorRep'), 60))
    hint(doc, '(наименование лица, осуществляющего строительство (генеральный строительный '
              'подрядчик), должность, Ф.И.О. представителя, дата, подпись)')

    sign_table(doc, [
        ('Предписание выдал ответственный представитель НСК:', 'Подпись, дата', False),
        (f"Инспектор СК {order.get('inspector') or ''}", '', False),
        ('(должность, ФИО)', '(подпись, дата)', True),
        ('Представитель Заказчика:', 'Подпись, дата', False),
        (line(body.get('customerRep'), 40), '', False),
        ('(должность, ФИО)', '(подпись, дата)', True),
    ])

    para(doc, 'Копии направлены:', space_before=10)
    para(doc, line(body.get('customerName'), 40))
    para(doc, line(issued_to, 40))

    para(doc,
         'Примечание: В случае отказа ответственного представителя строительного подрядчика '
         'от подписи, в Предписание в графу «Ответственный представитель лица, осуществляющего '
         'строительство» вносится запись «От подписи отказался».',
         size=9, space_before=10)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
